import os
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import crawl_public
from base import BaseCrawler


# 一次"正常结束"的抓取需要两个条件：解析出条目，且页面明确表示没有下一页
# （span.next 里没有 <a>）。只有末页标记而零条目会被判定为未完成——零条目
# 更可能意味着解析器失效，v1.52 的分页完整性保护有意为此保守处理。
LAST_PAGE_HTML = """
<html>
  <div class="item">
    <div class="title">
      <a href="https://movie.douban.com/subject/12345/">测试条目</a>
    </div>
    <span class="date">2026-01-01</span>
  </div>
  <span class="next"></span>
</html>
"""


class DummyResponse:
    def __init__(self, text=LAST_PAGE_HTML, status_code=200, url="https://example.test"):
        self.text = text
        self.status_code = status_code
        self.url = url


class PublicSessionTests(unittest.TestCase):
    def test_public_session_carries_no_cookies(self):
        session = crawl_public.build_public_session()

        self.assertEqual({}, session.cookies.get_dict())
        self.assertIn("Mozilla", session.headers["User-Agent"])

    def test_public_state_is_isolated_from_authenticated_state(self):
        """同一账号的公开数据和登录数据条目不同，基线不能共用。"""
        self.assertNotEqual("demo", crawl_public.state_key("demo"))
        self.assertIn("demo", crawl_public.state_key("demo"))


class PublicBackupFlowTests(unittest.TestCase):
    """公开模式应当走 base.py 的共享抓取栈，而不是另一套实现。"""

    def _run(self, tmpdir, requested, **kwargs):
        def fake_request(self, url, retries=3):
            requested.append(url)
            return DummyResponse(url=url)

        with patch.object(BaseCrawler, "_make_request", fake_request):
            return crawl_public.run_public_backup(
                "demo-user", output_dir=tmpdir, request_delay=0, **kwargs
            )

    def test_uses_shared_crawler_stack_and_correct_urls(self):
        requested = []
        with tempfile.TemporaryDirectory() as tmpdir:
            result = self._run(tmpdir, requested, categories=["books", "movies"])

        self.assertTrue(result["ok"])
        self.assertEqual({"books", "movies"}, set(result["data"]))
        # 书籍分页用 ?start= 而不是 &start=，否则豆瓣会忽略分页参数。
        self.assertIn(
            "https://book.douban.com/people/demo-user/collect?start=0&type=book",
            requested,
        )
        self.assertIn("https://movie.douban.com/people/demo-user/collect", requested)
        self.assertNotIn(
            "https://book.douban.com/people/demo-user/collect&start=0", requested
        )

    def test_exports_the_beautified_workbook(self):
        """公开模式改用 storage.py 之后，导出的是带元数据和总览页的报告。"""
        requested = []
        with tempfile.TemporaryDirectory() as tmpdir:
            self._run(tmpdir, requested, categories=["movies"])
            workbooks = [f for f in os.listdir(tmpdir) if f.endswith(".xlsx")]
            sheets = load_workbook(os.path.join(tmpdir, workbooks[0])).sheetnames

        self.assertEqual(1, len(workbooks))
        self.assertEqual("元数据", sheets[0])
        self.assertIn("总览", sheets)

    def test_metadata_records_public_mode(self):
        import json

        requested = []
        with tempfile.TemporaryDirectory() as tmpdir:
            self._run(tmpdir, requested, categories=["movies"])
            # 目录里还有 backup_baseline_*.json，必须精确匹配导出文件，
            # 否则取到哪个取决于 listdir 的任意顺序。
            dumps = [
                f
                for f in os.listdir(tmpdir)
                if f.startswith("douban_backup_") and f.endswith(".json")
            ]
            self.assertEqual(1, len(dumps))
            path = os.path.join(tmpdir, dumps[0])
            with open(path, encoding="utf-8") as file_obj:
                payload = json.load(file_obj)

        self.assertEqual("public", payload["metadata"]["backup_mode"])
        self.assertEqual("demo-user", payload["metadata"]["user_id"])

    def test_checkpoint_can_be_disabled(self):
        requested = []
        with tempfile.TemporaryDirectory() as tmpdir:
            self._run(
                tmpdir, requested, categories=["movies"], checkpoint_enabled=False
            )
            states = [f for f in os.listdir(tmpdir) if f.startswith("backup_state_")]

        self.assertEqual([], states)

    def test_completed_backup_writes_a_baseline(self):
        requested = []
        with tempfile.TemporaryDirectory() as tmpdir:
            self._run(tmpdir, requested, categories=["movies"], incremental=True)
            baselines = [
                f for f in os.listdir(tmpdir) if f.startswith("backup_baseline_")
            ]

        self.assertEqual(1, len(baselines))

    def test_incomplete_backup_leaves_no_baseline(self):
        """抓取中断时基线不能被残缺数据覆盖。"""
        requested = []

        def failing_request(self, url, retries=3):
            requested.append(url)
            return None  # 模拟请求彻底失败

        with tempfile.TemporaryDirectory() as tmpdir:
            with patch.object(BaseCrawler, "_make_request", failing_request):
                crawl_public.run_public_backup(
                    "demo-user",
                    categories=["movies"],
                    output_dir=tmpdir,
                    request_delay=0,
                    incremental=True,
                )
            baselines = [
                f for f in os.listdir(tmpdir) if f.startswith("backup_baseline_")
            ]

        self.assertEqual([], baselines)


class PublicCliTests(unittest.TestCase):
    def test_standalone_cli_passes_all_options(self):
        with patch.object(crawl_public, "run_public_backup") as run_public:
            crawl_public.main(
                ["demo-user", "--delay", "2.5", "--incremental", "--no-resume"]
            )

        run_public.assert_called_once_with(
            "demo-user",
            output_dir=None,
            request_delay=2.5,
            checkpoint_enabled=False,
            incremental=True,
            download_covers=False,
            full_reviews=False,
            formats=['xlsx'],
        )

    def test_standalone_cli_defaults(self):
        with patch.object(crawl_public, "run_public_backup") as run_public:
            crawl_public.main(["demo-user"])

        run_public.assert_called_once_with(
            "demo-user",
            output_dir=None,
            request_delay=None,
            checkpoint_enabled=True,
            incremental=False,
            download_covers=False,
            full_reviews=False,
            formats=['xlsx'],
        )

    def test_blank_user_id_aborts(self):
        with patch("builtins.input", return_value="  "), patch.object(
            crawl_public, "run_public_backup"
        ) as run_public:
            self.assertIsNone(crawl_public.main([]))

        run_public.assert_not_called()


if __name__ == "__main__":
    unittest.main()
